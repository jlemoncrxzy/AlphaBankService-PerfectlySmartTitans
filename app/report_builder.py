"""Сборка DOCX-отчёта по компании на основе загруженного шаблона."""
import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from sqlalchemy.orm import Session, joinedload

from .metrics import load_payment_stats
from .models import AnalysisRun, BankAnalytic, Client, EventInsight, NewsItem
from .routers.dashboard import format_client, format_insight

PLACEHOLDER_HINT = (
    "{{report_title}}, {{company_name}}, {{industry}}, {{revenue}}, "
    "{{region}}, {{products}}, {{date}}, {{incoming}}, {{outgoing}}, "
    "{{counterparties}}, {{risks_count}}, {{opportunities_count}}"
)

ILLEGAL_FILENAME = str.maketrans({char: "_" for char in '\\/:*?"<>|'})


def safe_filename(value: str, fallback: str = "company-report") -> str:
    cleaned = (value or "").translate(ILLEGAL_FILENAME).strip().strip(".")
    return cleaned[:120] or fallback


def _fmt_money(value) -> str:
    if value is None:
        return "—"
    try:
        amount = float(value)
    except (TypeError, ValueError):
        return "—"
    if amount <= 0:
        return "—"
    if amount >= 1e9:
        return f"{amount / 1e9:.1f} млрд ₽".replace(".", ",")
    if amount >= 1e6:
        return f"{amount / 1e6:.1f} млн ₽".replace(".", ",")
    if amount >= 1e3:
        return f"{amount / 1e3:.1f} тыс ₽".replace(".", ",")
    return f"{amount:.0f} ₽"


def _fmt_count(value) -> str:
    try:
        return f"{int(value):,}".replace(",", " ")
    except (TypeError, ValueError):
        return "—"


def _document_has_content(document: Document) -> bool:
    if any(paragraph.text.strip() for paragraph in document.paragraphs):
        return True
    if any(
        cell.text.strip()
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ):
        return True
    for section in document.sections:
        if any(paragraph.text.strip() for paragraph in section.header.paragraphs):
            return True
        if any(paragraph.text.strip() for paragraph in section.footer.paragraphs):
            return True
    return False


def _replace_in_paragraph(paragraph, mapping: dict[str, str]) -> int:
    original = "".join(run.text for run in paragraph.runs)
    if not original:
        return 0

    updated = original
    for token, value in mapping.items():
        updated = updated.replace(token, str(value))
    if updated == original:
        return 0

    if paragraph.runs:
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(updated)
    return 1


def _replace_in_table(table, mapping: dict[str, str]) -> int:
    replaced = 0
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replaced += _replace_in_paragraph(paragraph, mapping)
            for nested in cell.tables:
                replaced += _replace_in_table(nested, mapping)
    return replaced


def _replace_placeholders(document: Document, mapping: dict[str, str]) -> int:
    replaced = 0
    for paragraph in document.paragraphs:
        replaced += _replace_in_paragraph(paragraph, mapping)
    for table in document.tables:
        replaced += _replace_in_table(table, mapping)
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            replaced += _replace_in_paragraph(paragraph, mapping)
        for table in section.header.tables:
            replaced += _replace_in_table(table, mapping)
        for paragraph in section.footer.paragraphs:
            replaced += _replace_in_paragraph(paragraph, mapping)
        for table in section.footer.tables:
            replaced += _replace_in_table(table, mapping)
    return replaced


def _shade_cell(cell, color: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell_properties.append(shading)


def _set_table_borders(table, color: str = "D9D9DE") -> None:
    """Гарантирует аккуратные границы даже в шаблоне без стиля Table Grid."""
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)
    table._tbl.tblPr.append(borders)


def _add_heading(document: Document, text: str, level: int):
    try:
        return document.add_heading(text, level=level)
    except KeyError:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        run.bold = True
        run.font.size = Pt(16 if level == 1 else 13)
        return paragraph


def _add_note(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x6B, 0x6D, 0x75)


def _add_key_value_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=len(rows), cols=2)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    _set_table_borders(table)
    table.autofit = False

    for index, (label, value) in enumerate(rows):
        label_cell, value_cell = table.rows[index].cells
        label_cell.width = Cm(4.6)
        value_cell.width = Cm(12.4)
        label_cell.text = label
        value_cell.text = value
        for paragraph in label_cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
        _shade_cell(label_cell, "F2F2F4")

    document.add_paragraph()


def _add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style=None)
        paragraph.paragraph_format.left_indent = Cm(0.5)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run(f"• {item}")


def _add_numbered(document: Document, items: list[str]) -> None:
    for index, item in enumerate(items, start=1):
        paragraph = document.add_paragraph(style=None)
        paragraph.paragraph_format.left_indent = Cm(0.5)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run(f"{index}. {item}")


def _load_company_data(db: Session, client_id: int) -> tuple[dict, list[dict], AnalysisRun | None]:
    client = (
        db.query(Client)
        .options(joinedload(Client.profile))
        .filter(Client.id == client_id)
        .first()
    )
    if client is None:
        raise ValueError("Клиент не найден")

    latest_run = (
        db.query(AnalysisRun)
        .join(EventInsight, EventInsight.run_id == AnalysisRun.id)
        .filter(AnalysisRun.client_id == client.id)
        .order_by(AnalysisRun.created_at.desc(), AnalysisRun.id.desc())
        .first()
    )

    formatted: list[dict] = []
    if latest_run:
        insights = (
            db.query(EventInsight)
            .options(joinedload(EventInsight.run))
            .filter(EventInsight.run_id == latest_run.id)
            .order_by(EventInsight.id.desc())
            .all()
        )
        news_ids = [item.source_id for item in insights if item.source_type == "news"]
        analytic_ids = [item.source_id for item in insights if item.source_type == "analytic"]
        news_map = (
            {item.id: item for item in db.query(NewsItem).filter(NewsItem.id.in_(news_ids)).all()}
            if news_ids
            else {}
        )
        analytic_map = (
            {item.id: item for item in db.query(BankAnalytic).filter(BankAnalytic.id.in_(analytic_ids)).all()}
            if analytic_ids
            else {}
        )
        for insight in insights:
            source = (
                news_map.get(insight.source_id)
                if insight.source_type == "news"
                else analytic_map.get(insight.source_id)
            )
            formatted.append(format_insight(insight, source))

    company = format_client(client, len(formatted), load_payment_stats())
    return company, formatted, latest_run


def _placeholder_map(company: dict, insights: list[dict], title: str) -> dict[str, str]:
    payments = company.get("payments") or {}
    risks = sum(1 for item in insights if item.get("type") == "risk")
    opportunities = sum(1 for item in insights if item.get("type") == "opportunity")
    counterparties = (
        (payments.get("incomingCounterparties") or 0)
        + (payments.get("outgoingCounterparties") or 0)
    )
    mapping = {
        "{{report_title}}": title,
        "{{title}}": title,
        "{{company_name}}": company.get("name") or "—",
        "{{company}}": company.get("name") or "—",
        "{{industry}}": company.get("industry") or "—",
        "{{company_industry}}": company.get("industry") or "—",
        "{{revenue}}": _fmt_money(company.get("revenueRub")),
        "{{company_revenue}}": _fmt_money(company.get("revenueRub")),
        "{{region}}": company.get("region") or "—",
        "{{company_region}}": company.get("region") or "—",
        "{{products}}": ", ".join(company.get("products") or []) or "—",
        "{{company_products}}": ", ".join(company.get("products") or []) or "—",
        "{{date}}": datetime.now().strftime("%d.%m.%Y"),
        "{{incoming}}": _fmt_money(payments.get("incomingRub")),
        "{{outgoing}}": _fmt_money(payments.get("outgoingRub")),
        "{{counterparties}}": _fmt_count(counterparties) if payments.get("hasData") else "—",
        "{{risks_count}}": str(risks),
        "{{opportunities_count}}": str(opportunities),
        "{{insights_count}}": str(len(insights)),
    }
    return mapping


def _append_company_report(
    document: Document,
    company: dict,
    insights: list[dict],
    title: str,
    latest_run: AnalysisRun | None,
) -> None:
    payments = company.get("payments") or {}
    risks = [item for item in insights if item.get("type") == "risk"]
    opportunities = [item for item in insights if item.get("type") == "opportunity"]
    neutral = [item for item in insights if item.get("type") == "neutral"]

    _add_heading(document, title, 1)
    _add_note(
        document,
        f"Сформировано автоматически: {datetime.now().strftime('%d.%m.%Y %H:%M')}",
    )
    analysis_date = ""
    if latest_run:
        moment = latest_run.finished_at or latest_run.created_at
        if moment:
            analysis_date = f" Дата анализа: {moment.strftime('%d.%m.%Y %H:%M')}."
    summary = (
        f"Проанализировано {len(insights)} событий: рисков — {len(risks)}, "
        f"возможностей — {len(opportunities)}, нейтральных — {len(neutral)}."
        f"{analysis_date}"
    )
    paragraph = document.add_paragraph(summary)
    paragraph.paragraph_format.space_after = Pt(8)

    _add_heading(document, "1. Профиль компании", 2)
    _add_key_value_table(document, [
        ("Компания", company.get("name") or "—"),
        ("Отрасль", company.get("industry") or "—"),
        ("Выручка", _fmt_money(company.get("revenueRub"))),
        ("Регионы", company.get("region") or "—"),
        ("Продукты", ", ".join(company.get("products") or []) or "—"),
        ("Топ-контрагентов в профиле", _fmt_count(company.get("counterpartiesCount"))),
    ])

    _add_heading(document, "2. Финансовые показатели", 2)
    _add_key_value_table(document, [
        ("Поступления", _fmt_money(payments.get("incomingRub"))),
        ("Списания", _fmt_money(payments.get("outgoingRub"))),
        ("Контрагенты-плательщики", _fmt_count(payments.get("incomingCounterparties")) if payments.get("hasData") else "—"),
        ("Контрагенты-получатели", _fmt_count(payments.get("outgoingCounterparties")) if payments.get("hasData") else "—"),
        ("Входящих операций", _fmt_count(payments.get("incomingTransactions")) if payments.get("hasData") else "—"),
        ("Исходящих операций", _fmt_count(payments.get("outgoingTransactions")) if payments.get("hasData") else "—"),
        ("Рисков в ленте", str(len(risks))),
        ("Возможностей в ленте", str(len(opportunities))),
        ("Нейтральных событий", str(len(neutral))),
    ])

    _add_heading(document, "3. Риски и возможности", 2)
    if not insights:
        document.add_paragraph("Данные анализа по компании пока отсутствуют.")
    for index, item in enumerate(insights, start=1):
        _add_heading(document, f"3.{index}. {item.get('title') or 'Без названия'}", 3)
        kind = {
            "risk": "риск",
            "opportunity": "возможность",
            "neutral": "нейтральное событие",
        }.get(item.get("type"), "событие")
        confidence = item.get("confidence")
        meta = f"Тип: {kind}."
        if confidence is not None:
            meta += f" Уверенность: {confidence}%."
        if item.get("source"):
            meta += f" Источник: {item['source']}."
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run(meta)

        facts = item.get("details") or item.get("facts") or item.get("summary")
        if facts:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.add_run("Что произошло: ").bold = True
            paragraph.add_run(facts)

        impact = item.get("impact")
        if impact:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.add_run("Почему это важно: ").bold = True
            paragraph.add_run(impact)

        actions = item.get("actions") or []
        if actions:
            label = document.add_paragraph()
            label.paragraph_format.space_after = Pt(1)
            label.add_run("Что можно сделать:").bold = True
            _add_numbered(document, actions)

        document.add_paragraph()

    _add_heading(document, "4. Сводный план действий", 2)
    unique_actions: list[str] = []
    for item in insights:
        for action in item.get("actions") or []:
            if action not in unique_actions:
                unique_actions.append(action)
    if unique_actions:
        _add_bullets(document, unique_actions)
    else:
        document.add_paragraph("Рекомендации появятся после следующего запуска анализа.")

    _add_heading(document, "5. Источники", 2)
    if insights:
        _add_bullets(document, [
            f"{item.get('title')} — {item.get('source') or 'база знаний банка'} ({item.get('date') or '—'})"
            for item in insights
        ])
    else:
        document.add_paragraph("Источники не указаны.")

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Отчёт подготовлен AI-аналитиком Альфа-Бизнеса.")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x8F, 0x92, 0x99)


def build_company_report(
    db: Session,
    client_id: int,
    template_bytes: bytes,
    title: str,
) -> bytes:
    """Заполняет шаблон и добавляет к нему аккуратный отчёт по компании."""
    company, insights, latest_run = _load_company_data(db, client_id)
    document = Document(io.BytesIO(template_bytes))

    mapping = _placeholder_map(company, insights, title)
    _replace_placeholders(document, mapping)

    if _document_has_content(document):
        document.add_page_break()
    _append_company_report(document, company, insights, title, latest_run)

    document.core_properties.title = title
    document.core_properties.author = "Альфа-Бизнес"
    document.core_properties.subject = f"Отчёт по компании {company.get('name')}"

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
